"""Export sponsors.xlsx into dashboard/sponsors.json for the web dashboard to display."""
import json
import os

from openpyxl import load_workbook

XLSX_PATH = "sponsors.xlsx"
OUTPUT_PATH = os.path.join("dashboard", "sponsors.json")


def export():
    wb = load_workbook(XLSX_PATH)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    rows = [dict(zip(header, [c.value or "" for c in row])) for row in ws.iter_rows(min_row=2)]

    os.makedirs("dashboard", exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(rows, f, indent=2)
    print(f"Exported {len(rows)} row(s) to {OUTPUT_PATH}")


if __name__ == "__main__":
    export()
