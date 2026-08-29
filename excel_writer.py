"""Append sponsor rows from data/sponsors_cache.csv into a local Excel file, deduping by company+event."""
import csv
import datetime
import os

from openpyxl import Workbook, load_workbook

CACHE_PATH = os.path.join("data", "sponsors_cache.csv")
OUTPUT_PATH = os.environ.get("OUTPUT_XLSX_PATH", "sponsors.xlsx")
HEADER = ["Event Name", "Sponsor Company", "Contact Name", "Email", "Phone", "LinkedIn URL", "Source URL", "Date Found"]


def load_cache():
    if not os.path.exists(CACHE_PATH):
        return []
    with open(CACHE_PATH, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def get_workbook():
    if os.path.exists(OUTPUT_PATH):
        wb = load_workbook(OUTPUT_PATH)
        return wb, wb.active
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    return wb, ws


def sync():
    rows = load_cache()
    if not rows:
        print("No cached sponsor rows to sync.")
        return

    wb, ws = get_workbook()
    existing_keys = {(r[1].value, r[0].value) for r in ws.iter_rows(min_row=2) if r[0].value}

    today = datetime.date.today().isoformat()
    added = 0
    for row in rows:
        key = (row["company"], row["event_url"])
        if key in existing_keys:
            continue
        existing_keys.add(key)
        ws.append([
            row["event_url"], row["company"], row.get("contact_name", ""),
            row.get("email", ""), row.get("phone", ""), row.get("linkedin_url", ""),
            row.get("contact_source") or row["event_url"], today,
        ])
        added += 1

    wb.save(OUTPUT_PATH)
    print(f"Appended {added} new row(s) to {OUTPUT_PATH}")


if __name__ == "__main__":
    sync()
